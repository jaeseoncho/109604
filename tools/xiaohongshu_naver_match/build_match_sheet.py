from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
BLUE = "004E82"
kw = [
("에어컨·냉난방","에어컨 커버","空调防尘罩"),("에어컨·냉난방","실외기 커버","空调外机罩"),("에어컨·냉난방","에어컨 리모컨","空调遥控器"),
("에어컨·냉난방","에어컨 필터","空调过滤网"),("에어컨·냉난방","에어컨 배관 커버","空调管道装饰罩"),("에어컨·냉난방","에어컨 배수 호스","空调排水管"),
("에어컨·냉난방","에어컨 바람막이","空调挡风板"),("에어컨·냉난방","에어컨 청소 스프레이","空调清洗剂"),("에어컨·냉난방","에어컨 세척 커버","空调清洗罩"),
("에어컨·냉난방","리모컨 홀더","遥控器收纳架"),("에어컨·냉난방","실외기 받침대","空调外机支架"),("에어컨·냉난방","실외기 방진 패드","空调减震垫"),
("에어컨·냉난방","타이머 콘센트","定时插座"),("에어컨·냉난방","스마트 리모컨 허브","智能遥控器"),("에어컨·냉난방","온습도계","温湿度计"),
("에어컨·냉난방","서큘레이터","空气循环扇"),("에어컨·냉난방","탁상 선풍기","桌面小风扇"),("에어컨·냉난방","목걸이 선풍기","挂脖风扇"),
("에어컨·냉난방","휴대용 손선풍기","手持小风扇"),("에어컨·냉난방","냉풍기","冷风机"),("에어컨·냉난방","제습기","除湿机"),
("에어컨·냉난방","미니 제습기","迷你除湿器"),("에어컨·냉난방","이동식 에어컨","移动空调"),("에어컨·냉난방","창문형 에어컨","窗式空调"),
("에어컨·냉난방","에어컨 장식 스티커","空调装饰贴"),("에어컨·냉난방","창문 단열 필름","窗户隔热膜"),("에어컨·냉난방","암막 커튼","遮光窗帘"),
("에어컨·냉난방","문풍지","门窗密封条"),("에어컨·냉난방","에어컨 바람 가이드","空调导风板"),("에어컨·냉난방","실외기 차양막","空调外机遮阳罩"),
("에어컨·냉난방","전기 요금 측정기","电量计量插座"),("에어컨·냉난방","멀티탭","插线板"),("에어컨·냉난방","에어컨 전용 콘센트","空调专用插座"),
("에어컨·냉난방","사무실 무릎담요","办公室毯子"),("에어컨·냉난방","쿨매트","凉席"),("에어컨·냉난방","쿨링 베개 커버","凉感枕套"),
("에어컨·냉난방","냉감 이불","凉感被"),("에어컨·냉난방","아이스 넥밴드","冰凉颈环"),("에어컨·냉난방","에어컨 청소 브러시","空调清洁刷"),
("에어컨·냉난방","온풍기","暖风机"),
("난방·겨울","전기 히터","电暖器"),("난방·겨울","전기 장판","电热毯"),("난방·겨울","온수 매트","水暖毯"),("난방·겨울","발 난로","暖脚器"),
("난방·겨울","손난로","暖手宝"),("난방·겨울","가습기","加湿器"),("난방·겨울","미니 가습기","迷你加湿器"),("난방·겨울","단열 뽁뽁이","窗户保温膜"),
("난방·겨울","라디에이터 커버","暖气罩"),("난방·겨울","전기 방석","电热坐垫"),("난방·겨울","온도 조절기","温控器"),("난방·겨울","보일러 온도조절기","壁挂炉温控器"),
("난방·겨울","방풍 커튼","防风门帘"),("난방·겨울","핫팩","暖宝宝"),("난방·겨울","발열 조끼","发热马甲"),
("청소·관리","공기청정기 필터","空气净化器滤芯"),("청소·관리","공기청정기","空气净化器"),("청소·관리","차량용 공기청정기","车载空气净化器"),
("청소·관리","먼지 제거 브러시","除尘刷"),("청소·관리","틈새 청소 브러시","缝隙清洁刷"),("청소·관리","창틀 청소 도구","窗槽清洁刷"),
("청소·관리","곰팡이 제거제","除霉剂"),("청소·관리","배수구 냄새 차단","地漏防臭盖"),("청소·관리","리모컨 실리콘 케이스","遥控器保护套"),
("청소·관리","전선 정리 클립","理线器"),("청소·관리","벽걸이 수납 선반","壁挂置物架"),("청소·관리","무타공 선반","免打孔置物架"),
("청소·관리","무타공 후크","免打孔挂钩"),("청소·관리","벽 구멍 메꿈","墙洞修补膏"),("청소·관리","배관 구멍 마감 캡","空调孔装饰盖"),
("소형가전","미니 냉장고","迷你冰箱"),("소형가전","무선 청소기","无线吸尘器"),("소형가전","로봇 청소기","扫地机器人"),("소형가전","식기 건조대","碗碟沥水架"),
("소형가전","전기 포트","电热水壶"),("소형가전","에어프라이어","空气炸锅"),("소형가전","미니 세탁기","迷你洗衣机"),("소형가전","의류 건조기","干衣机"),
("소형가전","스팀 다리미","挂烫机"),("소형가전","전동 칫솔","电动牙刷"),("소형가전","무선 충전기","无线充电器"),("소형가전","보조배터리","充电宝"),
("소형가전","스마트 플러그","智能插座"),("소형가전","LED 무드등","氛围灯"),("소형가전","센서등","感应灯"),
("생활소품","밀폐용기","密封罐"),("생활소품","주방 수납 정리함","厨房收纳盒"),("생활소품","물병","水杯"),("생활소품","텀블러","保温杯"),
("생활소품","실리콘 주방도구","硅胶厨具"),("생활소품","수납 바구니","收纳篮"),("생활소품","옷걸이","衣架"),("생활소품","신발 정리대","鞋架"),
("생활소품","욕실 선반","浴室置物架"),("생활소품","샤워기 헤드","花洒"),("생활소품","도어 스토퍼","门挡"),("생활소품","창문 방충망","防蚊纱窗"),
("생활소품","모기 퇴치기","灭蚊灯"),("생활소품","캠핑 선풍기","露营风扇"),("생활소품","차량용 선풍기","车载风扇"),
]
assert len(kw) == 100, len(kw)

wb = Workbook()
ws = wb.active
ws.title = "매칭목록"
headers = ["번호","제품 분류","한국어 검색어","중국어 검색어","샤오홍슈 검색 링크","네이버쇼핑 검색 링크",
           "샤오홍슈 영상 URL","영상 내용 요약","네이버 상품 URL","네이버 상품명","판매가(원)",
           "일치 여부","사용 허가 확인","확인일","메모"]
widths = [6,13,20,18,20,22,42,30,42,30,12,11,16,12,30]
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin,right=thin,top=thin,bottom=thin)
hfill = PatternFill("solid", fgColor=BLUE)
yellow = PatternFill("solid", fgColor="FFF9C4")
for i,h in enumerate(headers,1):
    c = ws.cell(row=1,column=i,value=h)
    c.font = Font(name=FONT,bold=True,color="FFFFFF",size=11)
    c.fill = hfill; c.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
    ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
ws.row_dimensions[1].height = 32

for r,(cat,ko,zh) in enumerate(kw, start=2):
    ws.cell(row=r,column=1,value=r-1)
    ws.cell(row=r,column=2,value=cat)
    ws.cell(row=r,column=3,value=ko).font = Font(name=FONT,color="0000FF")
    ws.cell(row=r,column=4,value=zh).font = Font(name=FONT,color="0000FF")
    ws.cell(row=r,column=5,value=f'=HYPERLINK("https://www.xiaohongshu.com/search_result?keyword="&D{r},"샤오홍슈 검색 열기")')
    ws.cell(row=r,column=6,value=f'=HYPERLINK("https://search.shopping.naver.com/search/all?query="&C{r},"네이버쇼핑 검색 열기")')
    ws.cell(row=r,column=12,value="미확인")
    ws.cell(row=r,column=13,value="미확인")
    for col in (7,9):
        ws.cell(row=r,column=col).fill = yellow
    for col in range(1,16):
        c = ws.cell(row=r,column=col)
        c.border = border
        if c.font.color is None or c.font.color.rgb not in ("000000FF","FF0000FF"):
            c.font = Font(name=FONT, color=c.font.color) if c.font.color else Font(name=FONT)
        c.alignment = Alignment(vertical="center", wrap_text=(col in (8,10,15)))
    for col in (5,6):
        ws.cell(row=r,column=col).font = Font(name=FONT,color="0563C1",underline="single")
    ws.cell(row=r,column=11).number_format = '#,##0'
    ws.cell(row=r,column=14).number_format = 'yyyy-mm-dd'

dv1 = DataValidation(type="list", formula1='"일치,유사,불일치,미확인"', allow_blank=True,
                     error="목록에서 고르세요: 일치 / 유사 / 불일치 / 미확인", errorTitle="입력 오류", showErrorMessage=True)
dv2 = DataValidation(type="list", formula1='"원작자 허락받음,직접 촬영,미확인,사용 불가"', allow_blank=True,
                     error="목록에서 고르세요: 원작자 허락받음 / 직접 촬영 / 미확인 / 사용 불가", errorTitle="입력 오류", showErrorMessage=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add("L2:L101"); dv2.add("M2:M101")
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:O101"

# 요약 sheet
s = wb.create_sheet("요약")
s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 14
rows = [
 ("항목","값"),
 ("샤오홍슈 영상 URL 입력 수", '=COUNTA(매칭목록!G2:G101)'),
 ("네이버 상품 URL 입력 수", '=COUNTA(매칭목록!I2:I101)'),
 ("일치", '=COUNTIF(매칭목록!L2:L101,"일치")'),
 ("유사", '=COUNTIF(매칭목록!L2:L101,"유사")'),
 ("불일치", '=COUNTIF(매칭목록!L2:L101,"불일치")'),
 ("미확인", '=COUNTIF(매칭목록!L2:L101,"미확인")'),
 ("원작자 허락받음", '=COUNTIF(매칭목록!M2:M101,"원작자 허락받음")'),
 ("직접 촬영", '=COUNTIF(매칭목록!M2:M101,"직접 촬영")'),
 ("사용 가능 (일치 + 허락/직접촬영)", '=SUMPRODUCT((매칭목록!L2:L101="일치")*((매칭목록!M2:M101="원작자 허락받음")+(매칭목록!M2:M101="직접 촬영")))'),
 ("남은 작업(URL 미입력 행)", '=100-COUNTA(매칭목록!G2:G101)'),
]
for r,(a,b) in enumerate(rows, start=1):
    ca = s.cell(row=r,column=1,value=a); cb = s.cell(row=r,column=2,value=b)
    for c in (ca,cb):
        c.font = Font(name=FONT, bold=(r==1), color=("FFFFFF" if r==1 else None)); c.border=border
        if r==1: c.fill = hfill
    cb.alignment = Alignment(horizontal="right")
s.cell(row=13,column=1,value="※ 값은 '매칭목록' 시트를 채우면 자동으로 바뀝니다.").font = Font(name=FONT, italic=True, color="666666")

# 사용법 sheet
g = wb.create_sheet("사용법")
g.column_dimensions["A"].width = 110
lines = [
("샤오홍슈 제품 영상 ↔ 네이버쇼핑 판매 상품 매칭 정리표 (100개)", True),
("작성일: 2026-09-18", False),
("", False),
("[중요] 이 파일에는 영상 주소가 미리 들어 있지 않습니다.", True),
("샤오홍슈(xiaohongshu.com)와 네이버쇼핑은 자동 프로그램의 접속을 막고 있어, 영상 주소를 자동으로 모을 수 없었습니다.", False),
("확인되지 않은 주소를 지어내지 않기 위해, 검색어 100개와 검색 링크, 입력 칸을 준비했습니다. 영상 주소는 사용자가 직접 확인해서 붙여넣습니다.", False),
("", False),
("[색 안내]", True),
("파란 글씨(C·D열): 검색어. 원하는 제품으로 바꿔도 됩니다. 바꾸면 E·F열 검색 링크가 자동으로 따라 바뀝니다.", False),
("노란 칸(G·I열): 반드시 채워야 하는 칸. 샤오홍슈 영상 주소와 네이버 상품 주소.", False),
("E·F열: 클릭하면 브라우저에서 검색 결과가 열립니다. (수식이므로 지우지 마세요)", False),
("", False),
("[작업 순서]", True),
("1단계. '매칭목록' 시트에서 E열 '샤오홍슈 검색 열기'를 클릭합니다. (샤오홍슈는 로그인해야 영상이 보이는 경우가 많습니다)", False),
("2단계. 제품이 잘 보이는 영상 하나를 열고, 주소창의 주소를 복사해 G열에 붙여넣습니다.", False),
("        샤오홍슈 영상 주소 형식 예: https://www.xiaohongshu.com/explore/ 뒤에 영어·숫자로 된 고유번호", False),
("3단계. 영상에 나온 제품 특징(색, 모양, 기능)을 H열에 한 줄로 적습니다.", False),
("4단계. F열 '네이버쇼핑 검색 열기'를 클릭해 같은 제품을 찾고, 상품 주소를 I열, 상품명을 J열, 가격을 K열에 적습니다.", False),
("        네이버 상품 주소 형식 예: https://smartstore.naver.com/스토어이름/products/숫자", False),
("5단계. L열에서 일치 / 유사 / 불일치 중 하나를 고릅니다. (드롭다운)", False),
("6단계. M열에서 영상 사용 허가 상태를 고릅니다. 남의 영상을 허락 없이 내 클립에 올리면 저작권 문제로 삭제·제재될 수 있습니다.", False),
("7단계. N열에 확인한 날짜를 적습니다. '요약' 시트에서 진행 상황을 확인합니다.", False),
("", False),
("[주의]", True),
("- 샤오홍슈 영상은 원작자 허락 없이 그대로 재업로드하지 않습니다. 참고용(촬영 구도·구성 아이디어)으로만 쓰거나, 허락을 받은 뒤 사용합니다.", False),
("- 네이버 상품이 '유사'인 경우 영상 속 제품과 실제 판매 제품이 다를 수 있으므로 클립에 '동일 제품'이라고 쓰지 않습니다.", False),
("- 가격은 확인한 날짜 기준입니다. 시간이 지나면 달라질 수 있습니다.", False),
("- 검색어의 중국어 번역은 일반적인 표현으로 넣었습니다. 검색 결과가 적으면 D열을 다른 표현으로 바꿔 보세요.", False),
]
for r,(t,b) in enumerate(lines, start=1):
    c = g.cell(row=r,column=1,value=t)
    c.font = Font(name=FONT, bold=b, size=(14 if r==1 else 11), color=(BLUE if b else None))
    c.alignment = Alignment(wrap_text=True, vertical="top")

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save("xiaohongshu_naver_match_100.xlsx")
print("saved")
